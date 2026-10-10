"""Bounded, offline corpus parsing. This module never initializes cloud clients."""

from __future__ import annotations

import csv
import hashlib
import io
import json
import math
import re
import zipfile
from datetime import date, datetime
from pathlib import PurePath

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_CHUNKS = 100
MAX_TEXT_CHARS = 140_000
MAX_CHUNK_BYTES = 64 * 1024
MAX_METADATA_BYTES = 8 * 1024
SUPPORTED = {".json", ".jsonl", ".csv", ".xlsx", ".txt", ".md", ".markdown", ".pdf", ".docx"}
STRUCTURED = {".json", ".jsonl", ".csv", ".xlsx"}
MAPPING_FIELDS = (
    "content",
    "metadata",
    "law_name",
    "article_number",
    "source_url",
    "version",
    "effective_date",
    "checked_at",
    "case_number",
    "judgment_date",
    "cause",
    "jurisdiction",
)


class CorpusError(ValueError):
    def __init__(self, detail, code="invalid_corpus", status=422):
        super().__init__(detail)
        self.detail, self.code, self.status = detail, code, status


def canonical(value):
    return json.dumps(
        value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), allow_nan=False
    ).encode()


def digest(value):
    return hashlib.sha256(value if isinstance(value, bytes) else canonical(value)).hexdigest()


def text(value):
    return re.sub(
        r"\n{3,}", "\n\n", str(value or "").replace("\r\n", "\n").replace("\r", "\n")
    ).strip()


def metadata(value):
    if value is None:
        return {}
    if not isinstance(value, dict):
        raise CorpusError("metadata 必須是 JSON object")
    # Upload metadata must never author routing, ownership, or query scores.
    blocked = {
        "source_id",
        "collection",
        "collection_name",
        "distance",
        "retrieval_method",
        "embedding",
        "chunk_index",
        "chunk_count",
        "import_id",
    }

    def normalize(item, depth=0):
        if depth > 4:
            raise CorpusError("metadata 巢狀層數超過限制")
        if isinstance(item, (date, datetime)):
            return item.isoformat()
        if isinstance(item, dict):
            if len(item) > 64 or any(not isinstance(key, str) or len(key) > 128 for key in item):
                raise CorpusError("metadata 欄位數或欄位名稱超過限制")
            return {key: normalize(val, depth + 1) for key, val in item.items()}
        if isinstance(item, list):
            if len(item) > 32:
                raise CorpusError("metadata 陣列上限為 32 筆")
            return [normalize(val, depth + 1) for val in item]
        return item

    value = normalize(value)
    cleaned = {
        key: val for key, val in value.items() if key not in blocked and not key.startswith("_")
    }
    try:
        encoded = canonical(cleaned)
    except (ValueError, TypeError) as exc:
        raise CorpusError("metadata 包含不支援的值") from exc
    if len(encoded) > MAX_METADATA_BYTES:
        raise CorpusError("每筆 metadata 上限為 8 KiB")
    if "corpus_active" in cleaned and not isinstance(cleaned["corpus_active"], bool):
        raise CorpusError("corpus_active 必須是 boolean")
    return cleaned


def split_text(value, maximum=1600, overlap=180):
    """Preserve content with bounded overlapping windows, without silently truncating."""
    value = text(value)
    if len(value) <= maximum:
        return [value] if value else []
    chunks, start = [], 0
    while start < len(value):
        end = min(len(value), start + maximum)
        if end < len(value):
            boundary = value.rfind("\n", start + maximum // 2, end)
            if boundary > start:
                end = boundary
        chunks.append(value[start:end].strip())
        if end == len(value):
            break
        start = end - overlap
    return chunks


def _zip_guard(raw):
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            entries = archive.infolist()
            if len(entries) > 2000 or sum(item.file_size for item in entries) > 50 * 1024 * 1024:
                raise CorpusError("壓縮文件解壓後超過限制")
            if any(item.flag_bits & 1 for item in entries):
                raise CorpusError("不支援加密文件")
    except zipfile.BadZipFile as exc:
        raise CorpusError("文件不是有效的 ZIP 文件格式") from exc


def _headers(values):
    headers = [str(value).strip() if value is not None else "" for value in values]
    while headers and not headers[-1]:
        headers.pop()
    if not headers or any(not value or len(value) > 128 for value in headers):
        raise CorpusError("結構化欄名不可空白，且每個欄名上限為 128 字", "invalid_columns")
    if len(headers) > 128 or len(set(headers)) != len(headers):
        raise CorpusError("欄名不可重複，且上限為 128 欄", "invalid_columns")
    return headers


def _rows(raw, suffix):
    warnings = []
    if suffix == ".pdf":
        from pypdf import PdfReader, apply_configuration

        with apply_configuration(
            maximum_declared_stream_length=8 * 1024 * 1024,
            array_based_stream_maximum_output_length=8 * 1024 * 1024,
            zlib_maximum_output_length=8 * 1024 * 1024,
            lzw_maximum_output_length=8 * 1024 * 1024,
            run_length_maximum_output_length=8 * 1024 * 1024,
            brotli_maximum_output_length=8 * 1024 * 1024,
            stream_decoding_work_maximum_length=16 * 1024 * 1024,
            page_tree_maximum_entries=1000,
            xform_maximum_invocations_per_extraction=500,
            jbig2dec_binary=None,
        ):
            pdf = PdfReader(io.BytesIO(raw))
            if pdf.is_encrypted:
                raise CorpusError("不支援加密 PDF")
            if len(pdf.pages) > 100:
                raise CorpusError("PDF 上限為 100 頁")
            pages, total = [], 0
            for number, page in enumerate(pdf.pages, 1):
                extracted = text(page.extract_text())
                if not extracted and len(page.images):
                    raise CorpusError(f"PDF 第 {number} 頁需要 OCR；此版本不支援", "ocr_required")
                total += len(extracted)
                if total > MAX_TEXT_CHARS:
                    raise CorpusError("擷取文字超過 140,000 字，請拆分文件")
                pages.append(extracted)
        warnings.append("PDF 僅擷取文字；請核對下列全文的閱讀順序與表格。")
        return [{"content": "\n\n".join(pages)}], warnings
    if suffix == ".docx":
        from docx import Document

        _zip_guard(raw)
        document = Document(io.BytesIO(raw))
        # iter_inner_content preserves paragraph/table order.
        blocks = []
        for block in document.iter_inner_content():
            if hasattr(block, "rows"):
                blocks.extend("\t".join(cell.text for cell in row.cells) for row in block.rows)
            else:
                blocks.append(block.text)
        if document.inline_shapes:
            warnings.append("DOCX 內含圖片；圖片文字未經 OCR，不包含在下列擷取內容中。")
        return [{"content": "\n\n".join(blocks)}], warnings
    if suffix == ".xlsx":
        from openpyxl import load_workbook

        _zip_guard(raw)
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        rows = []
        try:
            for sheet in workbook:
                iterator = sheet.iter_rows(values_only=True)
                first = next(iterator, ())
                if not any(value is not None for value in first):
                    continue
                headers = _headers(first)
                for values in iterator:
                    if any(value is not None for value in values):
                        if any(value is not None for value in values[len(headers) :]):
                            raise CorpusError("資料包含沒有欄名的儲存格", "invalid_columns")
                        rows.append(dict(zip(headers, values, strict=False)))
                        if len(rows) > 100:
                            raise CorpusError("結構化資料上限為 100 筆")
        finally:
            workbook.close()
        return rows, warnings
    decoded = raw.decode("utf-8-sig")
    if suffix == ".json":
        value = json.loads(decoded)
        if isinstance(value, dict) and "documents" in value:
            value = value["documents"]
        return value if isinstance(value, list) else [value], warnings
    if suffix == ".jsonl":
        return [json.loads(line) for line in decoded.splitlines() if line.strip()], warnings
    if suffix == ".csv":
        reader = csv.reader(io.StringIO(decoded))
        headers = _headers(next(reader, ()))
        rows = []
        for values in reader:
            if not any(value for value in values):
                continue
            if len(values) > len(headers) and any(values[len(headers) :]):
                raise CorpusError("資料包含沒有欄名的儲存格", "invalid_columns")
            rows.append(dict(zip(headers, values, strict=False)))
            if len(rows) > 100:
                raise CorpusError("結構化資料上限為 100 筆")
        return rows, warnings
    frontmatter = {}
    if suffix in {".md", ".markdown"} and decoded.startswith("---\n"):
        end = decoded.find("\n---\n", 4)
        if end >= 0:
            header = decoded[4:end].strip()
            if header.startswith("{"):
                frontmatter = json.loads(header)
            else:
                for line in header.splitlines():
                    if ":" in line:
                        key, value = line.split(":", 1)
                        try:
                            frontmatter[key.strip()] = json.loads(value.strip())
                        except ValueError:
                            frontmatter[key.strip()] = value.strip().strip('"')
            decoded = decoded[end + 5 :]
    return [{"content": decoded, "metadata": frontmatter}], warnings


def validate_upload(filename, raw, kind, title=None, extra_metadata=None):
    """Only validate bounded upload metadata; never parse document contents."""
    if kind not in {"law", "judgment", "remedy"}:
        raise CorpusError("kind 必須為 law、judgment 或 remedy")
    filename = PurePath(str(filename).replace("\\", "/")).name
    if not filename or len(filename) > 240 or any(ord(char) < 32 for char in filename):
        raise CorpusError("檔名無效")
    suffix = PurePath(filename).suffix.lower()
    if suffix not in SUPPORTED:
        raise CorpusError("不支援此檔案格式")
    if not raw or len(raw) > MAX_FILE_BYTES:
        raise CorpusError("檔案不得為空，且不可超過 5 MiB")
    title = text(title or PurePath(filename).stem)
    if not title or len(title) > 300:
        raise CorpusError("來源名稱必須為 1 至 300 字")
    base_metadata = metadata(extra_metadata)
    return {"filename": filename, "suffix": suffix, "title": title, "metadata": base_metadata}


def validate_column_mapping(value, available_fields=None):
    if not isinstance(value, dict) or not value or set(value) - set(MAPPING_FIELDS):
        raise CorpusError("欄位對應只能使用支援的目標欄位", "invalid_mapping")
    if "content" not in value:
        raise CorpusError("請指定內容欄位", "invalid_mapping")
    for target, source in value.items():
        if not isinstance(source, str) or not source or len(source) > 128:
            raise CorpusError(f"{target} 必須對應一個有效的原始欄名", "invalid_mapping")
        if available_fields is not None and source not in available_fields:
            raise CorpusError(f"找不到原始欄位：{source}", "invalid_mapping")
    return dict(value)


def _read_rows(raw, suffix):
    try:
        rows, warnings = _rows(raw, suffix)
    except CorpusError:
        raise
    except Exception as exc:
        raise CorpusError(
            f"無法解析 {suffix} 文件；請檢查格式與 UTF-8 編碼", "parse_failed"
        ) from exc
    if not isinstance(rows, list) or not rows or len(rows) > 100:
        raise CorpusError("文件必須包含 1 至 100 筆文字記錄")
    if any(not isinstance(row, dict) for row in rows):
        raise CorpusError("每筆資料必須是 object，並含 content、text 或裁判書內容")
    return rows, warnings


def _available_fields(rows):
    fields = list(dict.fromkeys(key for row in rows for key in row))
    if len(fields) > 128 or any(
        not isinstance(key, str) or not key or len(key) > 128 for key in fields
    ):
        raise CorpusError("原始欄名必須為 1 至 128 字，且上限為 128 欄", "invalid_columns")
    return fields


def inspect_structured_upload(filename, raw):
    suffix = PurePath(filename).suffix.lower()
    if suffix not in STRUCTURED:
        raise CorpusError("此檔案不需要欄位對應", "mapping_not_supported")
    rows, warnings = _read_rows(raw, suffix)
    fields = _available_fields(rows)
    aliases = {
        "content": ("content", "text", "裁判書內容", "內容", "全文", "文件內容", "正文"),
        "metadata": ("metadata", "中繼資料"),
        "law_name": ("law_name", "法規名稱", "法律名稱"),
        "article_number": ("article_number", "條號"),
        "case_number": ("case_number", "裁判字號"),
        "judgment_date": ("judgment_date", "裁判日期"),
        "cause": ("cause", "裁判案由"),
    }
    suggested = {}
    for target in MAPPING_FIELDS:
        candidate = next(
            (field for field in aliases.get(target, (target,)) if field in fields), None
        )
        if candidate:
            suggested[target] = candidate
    return {
        "available_fields": fields,
        "suggested_mapping": suggested,
        "sample_rows": [{key: text(value)[:256] for key, value in row.items()} for row in rows[:3]],
        "row_count": len(rows),
        "warnings": warnings,
    }


def _map_rows(rows, column_mapping):
    mapping = validate_column_mapping(column_mapping, _available_fields(rows))
    mapped = []
    for number, row in enumerate(rows, 1):
        content = row.get(mapping["content"])
        if not isinstance(content, str) or not content.strip():
            raise CorpusError(f"第 {number} 筆的內容欄位必須為非空文字", "invalid_mapping")
        mapped.append(
            {
                target: row[source]
                for target, source in mapping.items()
                if row.get(source) is not None
            }
        )
    return mapped


def parse_upload(filename, raw, kind, title=None, extra_metadata=None, column_mapping=None):
    upload = validate_upload(filename, raw, kind, title, extra_metadata)
    filename, suffix, title = upload["filename"], upload["suffix"], upload["title"]
    base_metadata = upload["metadata"]
    rows, warnings = _read_rows(raw, suffix)
    if column_mapping is not None:
        if suffix not in STRUCTURED:
            raise CorpusError("此檔案不支援欄位對應", "mapping_not_supported")
        rows = _map_rows(rows, column_mapping)
    chunks, total_chars = [], 0
    for row in rows:
        if not isinstance(row, dict):
            raise CorpusError("每筆資料必須是 object，並含 content、text 或裁判書內容")
        content = text(row.get("content", row.get("text", row.get("裁判書內容", ""))))
        if not content:
            raise CorpusError("記錄沒有可擷取文字；掃描文件需要 OCR", "ocr_required")
        total_chars += len(content)
        if total_chars > MAX_TEXT_CHARS:
            raise CorpusError("擷取文字超過 140,000 字，請拆分文件")
        row_metadata = row.get("metadata", {})
        if isinstance(row_metadata, str):
            try:
                row_metadata = json.loads(row_metadata or "{}")
            except ValueError as exc:
                raise CorpusError("metadata 欄位必須為 JSON") from exc
        meta = {**base_metadata, **metadata(row_metadata)}
        for field in (
            "law_name",
            "article_number",
            "source_url",
            "version",
            "effective_date",
            "checked_at",
            "case_number",
            "judgment_date",
            "cause",
            "jurisdiction",
        ):
            if row.get(field) is not None:
                meta[field] = row[field]
        for chinese, field in (
            ("裁判字號", "case_number"),
            ("裁判日期", "judgment_date"),
            ("裁判案由", "cause"),
        ):
            if row.get(chinese) is not None:
                meta[field] = str(row[chinese])
        meta.update(
            data_type="document" if kind == "law" else kind, source=meta.get("source", title)
        )
        if meta.get("law_name") and meta.get("article_number"):
            from rag_core.rag.legal_lookup import lookup_key

            meta["lookup_keys"] = [lookup_key(str(meta["law_name"]), str(meta["article_number"]))]
            parts = [content]
        else:
            parts = split_text(content)
        for part in parts:
            if len(part.encode()) > MAX_CHUNK_BYTES:
                raise CorpusError("完整法條或 chunk 超過 64 KiB；不會自動截斷")
            chunks.append({"content": part, "metadata": metadata(meta)})
            if len(chunks) > MAX_CHUNKS:
                raise CorpusError("單次匯入超過 100 chunks，請拆分文件")
    for index, chunk in enumerate(chunks):
        chunk["metadata"].update(chunk_index=index, chunk_count=len(chunks))
    content_chunks = [
        {
            "content": chunk["content"],
            "metadata": {key: val for key, val in chunk["metadata"].items() if key != "source"},
        }
        for chunk in chunks
    ]
    return {
        "filename": filename,
        "title": title,
        "kind": kind,
        "chunks": chunks,
        "chunk_count": len(chunks),
        "warnings": warnings,
        "original_sha256": digest(raw),
        "content_hash": digest({"kind": kind, "chunks": content_chunks}),
        "identity_hash": digest({"kind": kind, "title": title.casefold()}),
    }


def validate_vector(vector, dimensions):
    if not isinstance(vector, list) or len(vector) != dimensions:
        raise CorpusError("Embedding 向量維度不符", "embedding_dimension_mismatch", 409)
    if any(
        isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value)
        for value in vector
    ):
        raise CorpusError("Embedding 向量包含無效數值", "invalid_embedding", 409)
    return [float(value) for value in vector]
