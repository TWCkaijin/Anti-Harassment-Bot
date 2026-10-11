"""Offline integration checks for source lifecycle and cross-database recovery."""

import copy
import hashlib
import io
import json
from types import SimpleNamespace

import pytest

from admin_app.corpus.parsing import CorpusError, parse_upload, validate_vector
from admin_app.corpus.service import IMPORTS, CorpusService
from admin_app.jobs import ContinueJob


class Snapshot:
    def __init__(self, ref, data):
        self.reference, self.id = ref, ref.id
        self.exists = data is not None
        self.data = copy.deepcopy(data)

    def to_dict(self):
        return copy.deepcopy(self.data)


class Reference:
    def __init__(self, db, path):
        self.db, self.path, self.id = db, path, path.rsplit("/", 1)[-1]

    def get(self, transaction=None):
        return Snapshot(self, self.db.data.get(self.path))

    def set(self, data):
        self.db.data[self.path] = copy.deepcopy(data)

    def create(self, data):
        if self.path in self.db.data:
            raise RuntimeError("already exists")
        self.set(data)

    def update(self, data):
        if self.db.fail_update_path == self.path:
            self.db.fail_update_path = None
            raise RuntimeError("private database temporarily unavailable")
        self.db.data[self.path].update(copy.deepcopy(data))

    def delete(self):
        self.db.data.pop(self.path, None)

    def collection(self, name):
        return Collection(self.db, f"{self.path}/{name}")


class Collection:
    def __init__(self, db, path):
        self.db, self.path = db, path
        self.maximum, self.after = None, None
        self.filter = None
        self.fields = None

    def document(self, name):
        return Reference(self.db, f"{self.path}/{name}")

    def order_by(self, field):
        assert field == "__name__"
        return self

    def select(self, fields):
        self.fields = fields
        return self

    def where(self, *, filter):
        self.filter = filter
        return self

    def limit(self, count):
        self.maximum = count
        return self

    def start_after(self, value):
        self.after = value["__name__"].id
        return self

    def stream(self):
        records = []
        for path in sorted(self.db.data):
            if path.startswith(self.path + "/") and "/" not in path[len(self.path) + 1 :]:
                ref = Reference(self.db, path)
                if self.filter is not None:
                    value = self.db.data[path]
                    for part in self.filter.field_path.split("."):
                        value = value.get(part, {})
                    if value != self.filter.value:
                        continue
                if self.after is None or ref.id > self.after:
                    snapshot = ref.get()
                    if self.fields is not None:
                        snapshot.data = {
                            key: value for key, value in snapshot.data.items() if key in self.fields
                        }
                    records.append(snapshot)
        return iter(records[: self.maximum] if self.maximum is not None else records)


class FakeDB:
    def __init__(self):
        self.data = {}
        self.fail_commit = False
        self.fail_update_path = None

    def collection(self, name):
        return Collection(self, name)

    def get_all(self, refs):
        return [ref.get() for ref in refs]


class Transaction:
    def __init__(self):
        self.operations = []

    def create(self, ref, value):
        self.operations.append(("create", ref.path, value))

    def set(self, ref, value):
        self.operations.append(("set", ref.path, value))

    def update(self, ref, value):
        self.operations.append(("update", ref.path, value))

    def delete(self, ref):
        self.operations.append(("delete", ref.path, None))


def atomic(db, callback):
    transaction = Transaction()
    result = callback(transaction)
    candidate = copy.deepcopy(db.data)
    for operation, path, value in transaction.operations:
        if operation == "create":
            assert path not in candidate
            candidate[path] = copy.deepcopy(value)
        elif operation == "set":
            candidate[path] = copy.deepcopy(value)
        elif operation == "update":
            candidate[path].update(copy.deepcopy(value))
        else:
            candidate.pop(path, None)
    if db.fail_commit:
        db.fail_commit = False
        raise RuntimeError("atomic commit failed")
    db.data = candidate
    return result


class Store:
    def __init__(self):
        self.objects = {}
        self.fail_delete_once = False

    def put_bytes(self, key, data, content_type):
        assert key not in self.objects
        self.objects[key] = data
        return {"name": key, "generation": "1", "bucket": "fake-private"}

    def read_bytes(self, ref):
        return self.objects[ref["name"]]

    def lookup_reference(self, key):
        return (
            {"name": key, "generation": "1", "bucket": "fake-private"}
            if key in self.objects
            else None
        )

    def put_bytes_idempotent(self, key, data, content_type):
        existing = self.lookup_reference(key)
        if existing:
            assert self.objects[key] == data
            return existing
        return self.put_bytes(key, data, content_type)

    def delete(self, ref):
        if self.fail_delete_once:
            self.fail_delete_once = False
            raise RuntimeError("storage temporarily unavailable")
        self.objects.pop(ref["name"], None)


class Context:
    def __init__(self):
        self.external = False
        self.checkpoints, self.updates = [], []

    def external_call_started(self, label):
        assert not self.external
        self.external = True

    def external_call_finished(self):
        self.external = False

    def checkpoint(self, value):
        self.checkpoints.append(value)

    def progress(self, value):
        self.updates.append(value)

    def ensure_active(self):
        pass


@pytest.fixture
def system(monkeypatch):
    business, private, store = FakeDB(), FakeDB(), Store()
    calls = []

    def embed(content):
        calls.append(content)
        return [0.1, 0.2, 0.3]

    service = CorpusService(business, private, store, atomic, embed)
    monkeypatch.setattr(
        CorpusService,
        "binding",
        staticmethod(
            lambda: {
                "project": "fake",
                "database": "app-dev",
                "provider": "fake",
                "model": "fake-e5",
                "dimensions": 3,
                "prefix_version": "e5-prefix-v1",
                "parser_version": 1,
            }
        ),
    )
    return SimpleNamespace(
        service=service, business=business, private=private, store=store, calls=calls
    )


@pytest.fixture
def queued(system, monkeypatch):
    import admin_app.jobs
    from admin_app.config import get_admin_settings

    system.queued_jobs = {}

    def submit(kind, payload, snapshot=None, idempotency_key=None):
        job_id = hashlib.sha256(f"{kind}:{idempotency_key}".encode()).hexdigest()[:32]
        if job_id not in system.queued_jobs:
            system.queued_jobs[job_id] = {
                "job_id": job_id,
                "kind": kind,
                "environment": get_admin_settings().environment,
                "status": "queued",
                "payload": payload,
                "result": None,
            }
        return system.queued_jobs[job_id]

    monkeypatch.setattr(admin_app.jobs, "submit_job", submit)
    monkeypatch.setattr(admin_app.jobs, "get_job", lambda job_id: system.queued_jobs[job_id])
    return system


def run_parse(system, draft):
    job = system.queued_jobs[draft["parse_job_id"]]
    job["result"] = system.service.parse_job(job, Context())
    job["status"] = "succeeded"
    return system.service.import_preview(draft["import_id"])


def import_job(preview, job_id="a" * 32):
    return {
        "job_id": job_id,
        "payload": {"import_id": preview["import_id"], "plan_hash": preview["plan_hash"]},
    }


def drain(system, job):
    context = Context()
    for _ in range(30):
        try:
            return system.service.import_job(job, context)
        except ContinueJob:
            assert not context.external
    raise AssertionError("job failed to complete bounded chunks")


def test_preview_has_complete_content_but_no_embedding_or_business_write(system):
    preview = system.service.preview("guide.md", "完整文字".encode(), "law")
    assert preview["chunks"][0]["content"] == "完整文字"
    assert not system.calls and not system.business.data
    assert len(system.store.objects) == 2


def test_atomic_publish_and_duplicate_worker_do_not_reembed(system):
    preview = system.service.preview("guide.txt", b"hello", "law")
    result = drain(system, import_job(preview))
    assert result["chunk_count"] == 1
    assert system.service.list_records("law")["items"][0]["content"] == "hello"
    assert system.service.source(result["source_id"])["status"] == "active"
    assert system.service.download(result["source_id"])[0] == b"hello"
    assert drain(system, import_job(preview)) == result
    assert len(system.calls) == 1


def test_failed_publish_is_atomic_and_retry_reuses_staging(system):
    preview = system.service.preview("guide.txt", b"hello", "law")
    system.business.fail_commit = True
    with pytest.raises(RuntimeError, match="atomic commit"):
        drain(system, import_job(preview))
    assert not system.business.data
    result = drain(system, import_job(preview))
    assert result["chunk_count"] == 1
    assert len(system.calls) == 1


def test_business_receipt_heals_private_status_failure(system):
    preview = system.service.preview("guide.txt", b"hello", "law")
    system.private.fail_update_path = f"{IMPORTS}/{preview['import_id']}"
    with pytest.raises(RuntimeError, match="private database"):
        drain(system, import_job(preview))
    assert system.service.list_records("law")["items"]
    result = drain(system, import_job(preview))
    assert result["source_id"] == preview["source_id"]
    assert len(system.calls) == 1


def test_large_import_yields_without_any_partial_publication(system):
    preview = system.service.preview("guide.txt", ("內容 " * 4000).encode(), "law")
    assert preview["chunk_count"] > 4
    with pytest.raises(ContinueJob):
        system.service.import_job(import_job(preview), Context())
    assert not system.service.list_records("law")["items"]
    assert len(system.calls) == 4
    result = drain(system, import_job(preview))
    assert result["chunk_count"] == preview["chunk_count"]
    assert len(system.calls) == preview["chunk_count"]


def test_stale_plan_or_changed_environment_fails_before_embedding(system, monkeypatch):
    preview = system.service.preview("guide.txt", b"hello", "law")
    job = import_job(preview)
    job["payload"]["plan_hash"] = "0" * 64
    with pytest.raises(CorpusError, match="任務"):
        drain(system, job)
    monkeypatch.setattr(CorpusService, "binding", staticmethod(lambda: {"database": "(default)"}))
    with pytest.raises(CorpusError, match="環境"):
        drain(system, import_job(preview))
    assert not system.calls


def test_dimension_mismatch_cannot_publish(system):
    preview = system.service.preview("guide.txt", b"hello", "law")
    system.service.embed = lambda content: [0.1, 0.2]
    with pytest.raises(CorpusError, match="維度"):
        drain(system, import_job(preview))
    assert not system.business.data


def test_source_conflict_detected_before_new_preview_or_embedding(system):
    preview = system.service.preview("guide.txt", b"hello", "law")
    drain(system, import_job(preview))
    with pytest.raises(CorpusError, match="相同來源"):
        system.service.preview("guide.txt", b"changed", "law")
    assert len(system.calls) == 1


def test_source_delete_retries_storage_cleanup_without_retrieval_reappearance(system):
    preview = system.service.preview("guide.txt", b"hello", "law")
    published = drain(system, import_job(preview))
    job = {"job_id": "b" * 32, "payload": {"source_id": published["source_id"], "revision": 1}}
    system.store.fail_delete_once = True
    with pytest.raises(RuntimeError, match="storage"):
        system.service.delete_job(job, Context())
    assert not system.service.list_records("law")["items"]
    assert system.service.source(published["source_id"])["status"] == "purge_pending"
    with pytest.raises(CorpusError, match="停止"):
        system.service.download(published["source_id"])
    result = system.service.delete_job(job, Context())
    assert result["original_deleted"] is True
    assert not system.store.objects
    assert system.service.delete_job(job, Context()) == result


def test_delete_stops_on_concurrent_content_change(system):
    preview = system.service.preview("guide.txt", b"hello", "law")
    published = drain(system, import_job(preview))
    ref = system.business.collection("rag_documents").document(f"{published['source_id']}_0000")
    ref.update({"content": "edited by another writer"})
    job = {"job_id": "b" * 32, "payload": {"source_id": published["source_id"], "revision": 1}}
    with pytest.raises(CorpusError, match="變更"):
        system.service.delete_job(job, Context())
    assert ref.get().exists
    assert system.store.objects


def test_delete_reservation_is_revision_fenced(system):
    preview = system.service.preview("guide.txt", b"hello", "law")
    published = drain(system, import_job(preview))
    job = {"job_id": "b" * 32, "payload": {"source_id": published["source_id"], "revision": 2}}
    with pytest.raises(CorpusError, match="版本"):
        system.service.delete_job(job, Context())
    assert system.service.source(published["source_id"])["status"] == "active"


def test_unknown_legacy_records_are_browsable_but_never_guessed_sources(system):
    system.business.collection("rag_documents").document("legacy").set(
        {"content": "whole old chunk", "metadata": {"source_file": "same-name.md"}}
    )
    assert system.service.list_records("law")["items"][0]["content"] == "whole old chunk"
    assert system.service.list_sources()["items"] == []


def test_record_pagination_is_stable_and_cursor_scope_is_checked(system):
    for index in range(5):
        system.business.collection("rag_documents").document(f"old{index}").set(
            {"content": str(index), "metadata": {}}
        )
    page = system.service.list_records("law", 2)
    page2 = system.service.list_records("law", 2, page["next_cursor"])
    assert [row["content"] for row in page["items"] + page2["items"]] == ["0", "1", "2", "3"]
    with pytest.raises(CorpusError):
        system.service.list_records("judgment", 2, page["next_cursor"])


@pytest.mark.parametrize(
    "filename,data",
    [
        ("data.json", b'{"documents":[{"content":"hello"}]}'),
        ("data.jsonl", b'{"text":"hello"}\n'),
        ("data.csv", b"content,source_url\nhello,https://example.org\n"),
        ("data.txt", b"hello"),
        ("data.md", b"# hello"),
    ],
)
def test_supported_text_formats(filename, data):
    parsed = parse_upload(filename, data, "law")
    assert "hello" in parsed["chunks"][0]["content"]


def test_complete_legal_article_and_untrusted_routing_metadata():
    body = "完整條文。" * 500
    parsed = parse_upload(
        "article.json",
        json.dumps(
            {
                "content": body,
                "metadata": {
                    "law_name": "性別平等工作法",
                    "article_number": "12",
                    "collection": "admin_credentials",
                    "source_id": "malicious",
                    "corpus_active": False,
                },
            }
        ).encode(),
        "law",
    )
    assert parsed["chunk_count"] == 1
    chunk = parsed["chunks"][0]
    assert chunk["content"] == body
    assert chunk["metadata"]["lookup_keys"] == ["性別平等工作法:12"]
    assert chunk["metadata"]["corpus_active"] is False
    assert "collection" not in chunk["metadata"] and "source_id" not in chunk["metadata"]


def test_limits_and_invalid_vector_fail_without_truncation():
    with pytest.raises(CorpusError, match="140,000"):
        parse_upload("large.txt", b"x" * 140001, "law")
    with pytest.raises(CorpusError, match="100"):
        parse_upload("rows.json", json.dumps([{"content": "x"}] * 101).encode(), "law")
    with pytest.raises(CorpusError, match="5 MiB"):
        parse_upload("large.txt", b"x" * (5 * 1024 * 1024 + 1), "law")
    for values in ([1, float("nan"), 0], [True, 1, 0]):
        with pytest.raises(CorpusError):
            validate_vector(values, 3)


def test_docx_preserves_paragraph_and_table_order():
    from docx import Document

    document = Document()
    document.add_paragraph("before")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "table"
    document.add_paragraph("after")
    output = io.BytesIO()
    document.save(output)
    content = parse_upload("file.docx", output.getvalue(), "law")["chunks"][0]["content"]
    assert content.index("before") < content.index("table") < content.index("after")


def test_xlsx_has_explicit_content_column():
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.active.append(["content", "case_number"])
    workbook.active.append(["judgment body", "case-1"])
    output = io.BytesIO()
    workbook.save(output)
    parsed = parse_upload("file.xlsx", output.getvalue(), "judgment")
    assert parsed["chunks"][0]["metadata"]["case_number"] == "case-1"


def test_pdf_extracts_text_and_enforces_page_limit():
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})}
    )
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 72 720 Td (hello PDF) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    output = io.BytesIO()
    writer.write(output)
    assert "hello PDF" in parse_upload("file.pdf", output.getvalue(), "law")["chunks"][0]["content"]
    for _ in range(100):
        writer.add_blank_page(width=612, height=792)
    output = io.BytesIO()
    writer.write(output)
    with pytest.raises(CorpusError, match="100 頁"):
        parse_upload("file.pdf", output.getvalue(), "law")


def legacy_source(system, count=201):
    for index in range(count):
        system.business.collection("rag_documents").document(f"legacy{index:04d}").set(
            {
                "content": f"old content {index}",
                "metadata": {
                    "source_file": "data/old.md",
                    "source": "old",
                    "chunk_index": index,
                    "chunk_count": count,
                },
            }
        )


def test_legacy_large_source_catalog_and_resumable_whole_source_delete(system):
    from admin_app.corpus.legacy import apply_legacy_catalog, prepare_legacy_catalog

    legacy_source(system)
    plan = prepare_legacy_catalog(system.service, "law", "data/old.md")
    assert not system.calls
    source = apply_legacy_catalog(system.service, plan, plan["plan_hash"])
    assert source["chunk_count"] == 201 and source["original_available"] is False
    page = system.service.source_records(source["source_id"], 100)
    assert len(page["chunks"]) == 100
    assert (
        len(system.service.source_records(source["source_id"], 100, page["next_cursor"])["chunks"])
        == 100
    )
    job = {"job_id": "d" * 32, "payload": {"source_id": source["source_id"], "revision": 1}}
    with pytest.raises(ContinueJob):
        system.service.delete_job(job, Context())
    assert len(list(system.business.collection("rag_documents").stream())) == 101
    with pytest.raises(ContinueJob):
        system.service.delete_job(job, Context())
    assert len(list(system.business.collection("rag_documents").stream())) == 1
    result = system.service.delete_job(job, Context())
    assert result["deleted_chunks"] == 201
    assert not system.store.objects


def test_legacy_incomplete_group_remains_unassociated(system):
    from admin_app.corpus.legacy import prepare_legacy_catalog

    legacy_source(system, 3)
    system.business.collection("rag_documents").document("legacy0001").delete()
    with pytest.raises(CorpusError, match="不完整"):
        prepare_legacy_catalog(system.service, "law", "data/old.md")
    assert system.service.list_sources()["items"] == []


def test_legacy_catalog_refuses_changed_snapshot(system):
    from admin_app.corpus.legacy import apply_legacy_catalog, prepare_legacy_catalog

    legacy_source(system, 3)
    plan = prepare_legacy_catalog(system.service, "law", "data/old.md")
    system.business.collection("rag_documents").document("legacy0001").update(
        {"content": "changed"}
    )
    with pytest.raises(CorpusError, match="計劃後"):
        apply_legacy_catalog(system.service, plan, plan["plan_hash"])
    assert not system.store.objects


def test_legacy_delete_refuses_new_outside_manifest_record(system):
    from admin_app.corpus.legacy import apply_legacy_catalog, prepare_legacy_catalog

    legacy_source(system, 3)
    plan = prepare_legacy_catalog(system.service, "law", "data/old.md")
    source = apply_legacy_catalog(system.service, plan, plan["plan_hash"])
    system.business.collection("rag_documents").document("late").set(
        {"content": "late addition", "metadata": {"source_file": "data/old.md"}}
    )
    job = {"job_id": "d" * 32, "payload": {"source_id": source["source_id"], "revision": 1}}
    with pytest.raises(CorpusError, match="清單外"):
        system.service.delete_job(job, Context())
    assert len(list(system.business.collection("rag_documents").stream())) == 4


def test_legacy_cleanup_retries_after_manifest_object_was_already_deleted(system):
    from admin_app.corpus.legacy import apply_legacy_catalog, prepare_legacy_catalog

    legacy_source(system, 3)
    plan = prepare_legacy_catalog(system.service, "law", "data/old.md")
    source = apply_legacy_catalog(system.service, plan, plan["plan_hash"])
    job = {"job_id": "d" * 32, "payload": {"source_id": source["source_id"], "revision": 1}}
    original_transaction = system.service.transaction
    calls = []

    def fail_final(db, callback):
        calls.append(1)
        if len(calls) == 3:
            raise RuntimeError("final business update failed")
        return original_transaction(db, callback)

    system.service.transaction = fail_final
    with pytest.raises(RuntimeError, match="final business"):
        system.service.delete_job(job, Context())
    assert not system.store.objects
    system.service.transaction = original_transaction
    assert system.service.delete_job(job, Context())["deleted_chunks"] == 3


def test_confirm_requires_exact_hash_and_same_job_idempotency(system, monkeypatch):
    import admin_app.jobs

    submitted = []

    def submit(kind, payload, snapshot=None, idempotency_key=None):
        submitted.append((kind, payload, snapshot, idempotency_key))
        return {"job_id": "a" * 32, "status": "queued"}

    monkeypatch.setattr(admin_app.jobs, "submit_job", submit)
    monkeypatch.setattr(
        admin_app.jobs, "get_job", lambda job_id: {"job_id": job_id, "status": "queued"}
    )
    preview = system.service.preview("guide.txt", b"hello", "law")
    with pytest.raises(CorpusError, match="預覽不符"):
        system.service.confirm(preview["import_id"], "wrong")
    assert not submitted and not system.calls
    assert system.service.confirm(preview["import_id"], preview["plan_hash"])["job_id"] == "a" * 32
    replay = system.service.confirm(preview["import_id"], preview["plan_hash"])
    assert replay["job_id"] == "a" * 32
    assert len(submitted) == 1


@pytest.mark.parametrize("saved_job_link", [True, False])
def test_confirm_after_publish_returns_same_job_and_preserves_original(
    system, monkeypatch, saved_job_link
):
    import admin_app.jobs

    preview = system.service.preview("synthetic.txt", b"synthetic content", "law")
    job = import_job(preview)
    published = drain(system, job)
    completed = {"job_id": job["job_id"], "status": "succeeded", "result": published}
    monkeypatch.setattr(admin_app.jobs, "get_job", lambda job_id: completed)
    if saved_job_link:
        system.private.collection(IMPORTS).document(preview["import_id"]).update(
            {"job_id": job["job_id"]}
        )
    before = copy.deepcopy(system.store.objects)
    replay = system.service.confirm(preview["import_id"], preview["plan_hash"])
    assert replay["job_id"] == job["job_id"]
    assert replay["result"] == published
    assert replay["status"] == "succeeded"
    assert system.store.objects == before
    assert system.service.download(published["source_id"])[0] == b"synthetic content"
    assert len(system.calls) == 1


def test_duplicate_cleanup_refuses_its_own_published_source(system):
    preview = system.service.preview("synthetic.txt", b"synthetic content", "law")
    published = drain(system, import_job(preview))
    data = system.service._import(preview["import_id"])
    with pytest.raises(CorpusError, match="已發布"):
        system.service._complete_duplicate(data, published["source_id"])
    assert system.service.download(published["source_id"])[0] == b"synthetic content"


def test_corpus_http_upload_preview_download_and_validation(system, queued, monkeypatch):
    from flask import Flask

    import admin_app.corpus

    monkeypatch.setattr(admin_app.corpus, "CorpusService", lambda: system.service)
    app = Flask(__name__)
    app.register_blueprint(admin_app.corpus.corpus_bp, url_prefix="/corpus")
    client = app.test_client()
    response = client.post(
        "/corpus/imports/preview",
        data={"file": (io.BytesIO(b"hello"), "file.txt"), "kind": "law"},
        content_type="multipart/form-data",
    )
    assert response.status_code == 202
    assert response.headers["Cache-Control"] == "no-store"
    assert not system.calls
    draft = response.get_json()
    assert draft["status"] == "parse_queued"
    preview = run_parse(system, draft)
    result = drain(system, import_job(preview))
    download = client.get(f"/corpus/sources/{result['source_id']}/download")
    assert download.data == b"hello"
    assert "attachment" in download.headers["Content-Disposition"]
    assert client.get("/corpus/records?kind=admin_credentials").status_code == 400
    assert client.get("/corpus/records?limit=oops").status_code == 400
    assert (
        client.post(f"/corpus/imports/{preview['import_id']}/confirm", json={}).status_code == 422
    )
    assert client.delete(f"/corpus/sources/{result['source_id']}", json={}).status_code == 422


def test_duplicate_upload_confirmation_cleans_only_its_own_original(system):
    first = system.service.preview("one.txt", b"same text", "law")
    source = drain(system, import_job(first))
    initial = set(system.store.objects)
    duplicate = system.service.preview("renamed.txt", b"same text", "law")
    assert duplicate["duplicate_source_id"] == source["source_id"]
    result = system.service.confirm(duplicate["import_id"], duplicate["plan_hash"])
    assert result["duplicate"] is True
    assert set(system.store.objects) == initial
    assert len(system.calls) == 1
    assert system.service.confirm(duplicate["import_id"], duplicate["plan_hash"]) == result
    assert system.service.download(source["source_id"])[0] == b"same text"


def test_duplicate_worker_race_publishes_one_source_and_cleans_loser_assets(system):
    first = system.service.preview("one.txt", b"same text", "law")
    second = system.service.preview("renamed.txt", b"same text", "law")
    winner = drain(system, import_job(first))
    loser = drain(system, import_job(second, "f" * 32))
    assert loser["source_id"] == winner["source_id"] and loser["duplicate"] is True
    assert len(system.service.list_sources()["items"]) == 1
    assert len(system.store.objects) == 2


def test_legacy_deletion_reduces_rejected_atomic_batch_size(system):
    from google.api_core.exceptions import InvalidArgument

    from admin_app.corpus.legacy import apply_legacy_catalog, prepare_legacy_catalog

    legacy_source(system, 20)
    plan = prepare_legacy_catalog(system.service, "law", "data/old.md")
    source = apply_legacy_catalog(system.service, plan, plan["plan_hash"])
    original_atomic = system.service.transaction

    def bounded_atomic(db, callback):
        transaction = Transaction()
        callback(transaction)
        if sum(operation[0] == "delete" for operation in transaction.operations) > 5:
            raise InvalidArgument("Transaction too big")
        return original_atomic(db, callback)

    system.service.transaction = bounded_atomic
    job = {"job_id": "d" * 32, "payload": {"source_id": source["source_id"], "revision": 1}}
    with pytest.raises(ContinueJob):
        system.service.delete_job(job, Context())
    assert len(list(system.business.collection("rag_documents").stream())) == 15
